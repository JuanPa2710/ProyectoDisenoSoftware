from tempfile import TemporaryDirectory
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import sys
import csv
import argparse
from pyspark.sql.types import ArrayType, MapType, StructType
from pathlib import Path
from pyspark.sql import SparkSession
from pyspark.sql import functions as F


def rechazar_archivo(codigo, mensaje):
    print(f"SIVCR_VALIDACION:{codigo}", file=sys.stderr, flush=True)
    print(mensaje, file=sys.stderr, flush=True)
    raise SystemExit(2)

parser = argparse.ArgumentParser(
    description="Procesar accidentes de COSEVI con Spark."
)
parser.add_argument(
    "--entrada",
    required=True,
    type=Path,
    help="Ruta del archivo CSV, JSON o Excel (.xlsx)  de accidentes."
)
parser.add_argument(
    "--salida",
    type=Path,
    help="Carpeta nueva donde guardar los JSON procesados."
)
argumentos = parser.parse_args()
ruta_entrada = argumentos.entrada.expanduser().resolve()

if not ruta_entrada.is_file():
    parser.error(f"No se encontró el archivo: {ruta_entrada}")

# Columnas necesarias para transformar el perfil de accidentes COSEVI.
columnas_requeridas = [
    "Clase de accidente",
    "Tipo de accidente",
    "Año",
    "Hora",
    "Hora recodificada",
    "Provincia",
    "Cantón",
    "Distrito",
    "Ruta",
    "Kilómetro",
    "Rural o urbano",
    "Calzada vertical",
    "Calzada horizontal",
    "Tipo de calzada",
    "Tipo de circulación",
    "Estado del tiempo",
    "Estado de la calzada",
    "Región Mideplan",
    "Tipo ruta",
    "Día",
    "Mes",
]

def validar_columnas(nombres):
    nombres = [nombre.strip() for nombre in nombres]

    if not nombres or any(not nombre for nombre in nombres):
        rechazar_archivo(
            "COLUMNAS_INCORRECTAS",
            "El archivo no contiene nombres de columnas válidos."
        )

    comparables = [nombre.casefold() for nombre in nombres]

    if len(comparables) != len(set(comparables)):
        rechazar_archivo(
            "COLUMNAS_INCORRECTAS",
            "Hay nombres de columnas repetidos."
        )

    faltantes = [
        nombre for nombre in columnas_requeridas
        if nombre not in nombres
    ]

    if faltantes:
        rechazar_archivo(
            "COLUMNAS_INCORRECTAS",
            "Faltan columnas del perfil COSEVI: "
            + ", ".join(faltantes)
        )


def convertir_excel_a_csv(origen, destino):
    try:
        libro = load_workbook(
            origen,
            read_only=True,
            data_only=False,
            keep_links=False
        )
    except (BadZipFile, InvalidFileException, ParseError, OSError, ValueError):
        rechazar_archivo(
            "EXCEL_INVALIDO",
            "No se pudo abrir el archivo Excel."
        )

    try:
        if len(libro.worksheets) != 1:
            rechazar_archivo(
                "EXCEL_INVALIDO",
                "El Excel debe contener exactamente una hoja."
            )

        hoja = libro.worksheets[0]
        filas = hoja.iter_rows()
        primera_fila = next(filas, None)

        if primera_fila is None:
            rechazar_archivo(
                "SIN_REGISTROS",
                "La hoja de Excel está vacía."
            )

        if any(celda.data_type in {"f", "e"} for celda in primera_fila):
            rechazar_archivo(
                "EXCEL_INVALIDO",
                "Los encabezados no pueden contener fórmulas ni errores."
            )

        encabezados = [
            "" if celda.value is None else str(celda.value).strip()
            for celda in primera_fila
        ]

        validar_columnas(encabezados)
        cantidad = 0

        with destino.open("w", encoding="utf-8", newline="") as salida:
            escritor = csv.writer(salida, delimiter=";")
            escritor.writerow(encabezados)

            for numero_fila, fila in enumerate(filas, start=2):
                valores = [celda.value for celda in fila]

                # Ignorar filas completamente vacías.
                if all(
                    valor is None
                    or (isinstance(valor, str) and not valor.strip())
                    for valor in valores
                ):
                    continue

                if any(celda.data_type in {"f", "e"} for celda in fila):
                    rechazar_archivo(
                        "EXCEL_INVALIDO",
                        f"La fila {numero_fila} contiene fórmulas o errores. "
                        "Exportá el archivo con valores."
                    )

                escritor.writerow([
                    "" if valor is None else str(valor)
                    for valor in valores
                ])
                cantidad += 1

        if cantidad == 0:
            rechazar_archivo(
                "SIN_REGISTROS",
                "El Excel contiene encabezados, pero no tiene registros."
            )
    finally:
        libro.close()

formato = ruta_entrada.suffix.lower()

if formato not in {".csv", ".json",".xlsx"}:
    parser.error("Se permiten archivos CSV, JSON y ·Excel (.xlsx).")

# Revisar el encabezado original del CSV antes de que Spark lo lea.
if formato == ".csv":
    try:
        with ruta_entrada.open(
            encoding="utf-8-sig", newline=""
        ) as entrada:
            lector = csv.reader(entrada, delimiter=";", strict=True)
            encabezado = next(lector, None)
    except (UnicodeError, csv.Error):
        parser.error(
            "Se requiere un CSV UTF-8 separado por punto y coma."
        )
    except OSError:
        parser.error("No se pudo abrir el archivo de entrada.")

    if not encabezado:
        rechazar_archivo(
            "SIN_REGISTROS",
            "El archivo está vacío o no contiene encabezados."
        )

    validar_columnas(encabezado)

spark = (
    SparkSession.builder
    .appName("CargaAccidentesSIVCR")
    .master("local[2]")
    .config("spark.ui.enabled", "false")
    .getOrCreate()
)

spark.sparkContext.setLogLevel("ERROR")

temporal = None

try:
    archivo = str(ruta_entrada)

    if formato == ".xlsx":
        temporal = TemporaryDirectory(prefix="sivcr-excel-")
        archivo_temporal = Path(temporal.name) / "entrada.csv"
        convertir_excel_a_csv(ruta_entrada, archivo_temporal)
        archivo = str(archivo_temporal)

    if formato in {".csv", ".xlsx"}:
        datos = (
            spark.read
            .option("header", "true")
            .option("sep", ";")
            .option("encoding", "UTF-8")
            .csv(archivo)
        )
    else:
        datos = (
            spark.read
            .option("multiLine", "true")
            .option("mode", "FAILFAST")
            .option("primitivesAsString", "true")
            .json(archivo)
        )

    if not datos.take(1):
        rechazar_archivo(
            "SIN_REGISTROS",
            "El archivo no contiene registros para procesar."
        )

    validar_columnas(datos.columns)

    # Quitar espacios al inicio y al final de los encabezados.
    datos = datos.toDF(*[nombre.strip() for nombre in datos.columns])



    # Trabajar únicamente con las columnas de este perfil.
    datos = datos.select(*columnas_requeridas)


    # Este perfil espera campos simples, no objetos ni listas anidadas.
    campos_complejos = [
        campo.name
        for campo in datos.schema.fields
        if isinstance(campo.dataType, (ArrayType, MapType, StructType))
    ]

    if campos_complejos:
        rechazar_archivo(
            "COLUMNAS_INCORRECTAS",
            "Estos campos deben contener valores simples: "
            + ", ".join(campos_complejos)
        )


    # Quitar espacios al inicio y al final de cada valor.
    datos = datos.select(
        *[F.trim(F.col(nombre)).alias(nombre) for nombre in datos.columns]
    )

    print("\nCOLUMNAS LIMPIAS:", datos.columns)
    print("\nTOTAL DE REGISTROS:", datos.count())

    datos.select(
        "Año", "Provincia", "Cantón", "Día", "Mes"
    ).show(5, truncate=False)
    # Revisar los valores y su frecuencia antes de transformarlos.
    for campo in ["Año", "Mes", "Día", "Kilómetro"]:
        print(f"\nVALORES DE: {campo}")

        datos.groupBy(campo).count().orderBy(
            F.desc("count")
        ).show(15, truncate=False)


    # Conservar Kilómetro y crear una versión numérica aparte.
    datos = datos.withColumn(
        "kilometro_numero",
        F.expr("try_cast(replace(`Kilómetro`, ',', '.') AS DOUBLE)")
    )

    print("\nEJEMPLOS DE CONVERSIÓN:")
    datos.select(
        "Kilómetro", "kilometro_numero"
    ).distinct().orderBy("Kilómetro").show(20, truncate=False)

    print("\nVALORES QUE NO SE PUDIERON CONVERTIR:")
    datos.filter(
        F.col("kilometro_numero").isNull()
    ).groupBy("Kilómetro").count().orderBy(
        F.desc("count")
    ).show(30, truncate=False)

    # Convertir el año a un número entero.
    datos = datos.withColumn(
        "anio",
        F.expr("try_cast(`Año` AS INT)")
    )

    # Relacionar las etiquetas del archivo con números de mes.
    meses = {
        "A. Enero": 1,
        "B. Febrero": 2,
        "C. Marzo": 3,
        "D. Abril": 4,
        "E. Mayo": 5,
        "F. Junio": 6,
        "G. Julio": 7,
        "H. Agosto": 8,
        "I. Setiembre": 9,
        "J. Octubre": 10,
        "K. Noviembre": 11,
        "L. Diciembre": 12,
    }

    pares = []
    for etiqueta, numero in meses.items():
        pares.extend([F.lit(etiqueta), F.lit(numero)])

    mapa_meses = F.create_map(*pares)

    datos = datos.withColumn(
        "mes_numero",
        F.element_at(mapa_meses, F.col("Mes"))
    )

    print("\nAÑO Y MES CONVERTIDOS:")
    datos.select(
        "Año", "anio", "Mes", "mes_numero"
    ).show(5, truncate=False)

    print("\nREGISTROS SIN AÑO O MES NUMÉRICO:")
    print(
        datos.filter(
            F.col("anio").isNull()
            | F.col("mes_numero").isNull()
        ).count()
    )

    # Traducir las etiquetas de COSEVI a números de día de semana.
    dias = {
        "1.Domingo": 1,
        "2.Lunes": 2,
        "3.Martes": 3,
        "4.Miércoles": 4,
        "5.Jueves": 5,
        "6.Viernes": 6,
        "7.Sábado": 7,
    }

    pares_dias = []
    for etiqueta, numero in dias.items():
        pares_dias.extend([F.lit(etiqueta), F.lit(numero)])

    mapa_dias = F.create_map(*pares_dias)

    datos = datos.withColumn(
        "dia_semana_numero",
        F.element_at(mapa_dias, F.col("Día"))
    )

    print("\nDÍAS CONVERTIDOS:")
    datos.select(
        "Día", "dia_semana_numero"
    ).distinct().orderBy("dia_semana_numero").show(
        10, truncate=False
    )

    print("\nREGISTROS SIN DÍA RECONOCIDO:")
    print(
        datos.filter(
            F.col("dia_semana_numero").isNull()
        ).count()
    )

    # Validaciones ordenadas: conservar solo el primer fallo.
    from datetime import datetime

    anio_actual = datetime.now().year

    def fallo(condicion, campo, mensaje):
        return F.when(
            condicion,
            F.struct(
                F.lit(campo).alias("campo"),
                F.col(campo).cast("string").alias("valor"),
                F.lit(mensaje).alias("problema")
            )
        )

    def esta_vacio(campo):
        return (
            F.col(campo).isNull()
            | (F.trim(F.col(campo)) == "")
        )

    provincias_validas = [
        "san josé",
        "alajuela",
        "cartago",
        "heredia",
        "guanacaste",
        "puntarenas",
        "limón",
    ]

    datos = datos.withColumn(
        "primer_error",
        F.coalesce(
            fallo(
                F.col("anio").isNull()
                | ~F.col("anio").between(1900, anio_actual),
                "Año",
                f"Debe ser un año entero entre 1900 y {anio_actual}."
            ),
            fallo(
                F.col("mes_numero").isNull(),
                "Mes",
                "El mes no pertenece al catálogo del perfil."
            ),
            fallo(
                F.col("dia_semana_numero").isNull(),
                "Día",
                "El día de la semana no pertenece al catálogo."
            ),
            fallo(
                esta_vacio("Provincia")
                | ~F.lower(F.col("Provincia")).isin(provincias_validas),
                "Provincia",
                "Debe indicar una de las siete provincias de Costa Rica."
            ),
            fallo(
                esta_vacio("Clase de accidente"),
                "Clase de accidente",
                "La clase de accidente es obligatoria."
            ),
            fallo(
                esta_vacio("Tipo de accidente"),
                "Tipo de accidente",
                "El tipo de accidente es obligatorio."
            )
        )
    )

    print("\nRESUMEN DE VALIDACIÓN POR REGISTRO:")
    datos.select(
        F.when(
            F.col("primer_error").isNull(),
            F.lit("Válido")
        ).otherwise(F.lit("Rechazado")).alias("resultado")
    ).groupBy("resultado").count().show(truncate=False)

    print("\nPRIMER FALLO DE LOS REGISTROS RECHAZADOS:")
    datos.filter(
        F.col("primer_error").isNotNull()
    ).select(
        "primer_error.campo",
        "primer_error.valor",
        "primer_error.problema"
    ).show(20, truncate=False)

    # Separar los registros sin modificar el conjunto original.
    validos = datos.filter(F.col("primer_error").isNull())
    rechazados = datos.filter(F.col("primer_error").isNotNull())

    total_leidos = datos.count()
    total_aceptados = validos.count()
    total_rechazados = rechazados.count()

    if total_leidos != total_aceptados + total_rechazados:
        raise RuntimeError("Los conteos de validación no coinciden.")


    # Organizar los datos en documentos con campos agrupados.
    documentos = validos.select(
        F.lit("COSEVI").alias("fuente"),

        F.col("Clase de accidente").alias("clase_accidente"),
        F.col("Tipo de accidente").alias("tipo_accidente"),

        F.struct(
            F.col("anio"),
            F.col("mes_numero").alias("mes"),
            F.col("dia_semana_numero").alias("dia_semana"),
            F.col("Hora").alias("intervalo_hora")
        ).alias("tiempo"),

        F.struct(
            F.col("Provincia").alias("provincia"),
            F.col("Cantón").alias("canton"),
            F.col("Distrito").alias("distrito"),
            F.col("Ruta").alias("ruta"),
            F.col("kilometro_numero").alias("kilometro"),
            F.col("Kilómetro").alias("kilometro_original")
                ).alias("ubicacion"),

        F.struct(
            F.col("Estado del tiempo").alias("clima"),
            F.col("Estado de la calzada").alias("estado_calzada"),
            F.col("Tipo de calzada").alias("tipo_calzada"),
            F.col("Calzada vertical").alias("calzada_vertical"),
            F.col("Calzada horizontal").alias("calzada_horizontal"),
            F.col("Tipo de circulación").alias("tipo_circulacion"),
            F.col("Rural o urbano").alias("zona")
        ).alias("condiciones"),

        F.col("Región Mideplan").alias("region_mideplan"),
        F.col("Tipo ruta").alias("tipo_ruta"),
        F.col("Hora recodificada").alias("franja_horaria")
    )

    import json
    from datetime import datetime

    # Mostrar un ejemplo solo si existe algún documento válido.
    print("\nEJEMPLO DE DOCUMENTO JSON:")
    ejemplos = documentos.limit(1).toJSON(
        use_unicode=True
    ).collect()

    if ejemplos:
        print(json.dumps(
            json.loads(ejemplos[0]),
            indent=2,
            ensure_ascii=False
        ))
    else:
        print("No hay registros válidos para importar.")

    # Crear una carpeta distinta para cada ejecución.
    if argumentos.salida is not None:
        salida = str(argumentos.salida.expanduser().resolve())
    else:
        nombre_carga = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        salida = str(
            Path.home() / "apps/datos/procesados" / nombre_carga
        )

    # En la raíz se guardan únicamente documentos válidos.
    (
        documentos.write
        .mode("errorifexists")
        .option("ignoreNullFields", "false")
        .json(salida)
    )

    # Usar el esquema conocido permite comprobar también una salida vacía.
    guardados = (
        spark.read
        .schema(documentos.schema)
        .json(str(Path(salida) / "part-*.json"))
        .count()
    )

    if guardados != total_aceptados:
        raise RuntimeError(
            "La cantidad guardada no coincide con los registros aceptados."
        )

    # Conservar cada registro rechazado junto con su primer fallo.
    detalle_rechazados = rechazados.select(
        F.col("primer_error.campo").alias("campo"),
        F.col("primer_error.valor").alias("valor"),
        F.col("primer_error.problema").alias("problema"),
        F.struct(
            *[F.col(nombre) for nombre in columnas_requeridas]
        ).alias("registro_original")
    )

    carpeta_rechazados = Path(salida) / "rechazados"

    if total_rechazados > 0:
        (
            detalle_rechazados.write
            .mode("errorifexists")
            .option("ignoreNullFields", "false")
            .json(str(carpeta_rechazados))
        )

    # Resumen por primer campo que falló.
    errores_por_campo = {
        fila["campo"]: fila["count"]
        for fila in (
            detalle_rechazados
            .groupBy("campo")
            .count()
            .collect()
        )
    }

    # Una muestra pequeña para mostrar posteriormente en la pantalla.
    muestra_errores = [
        fila.asDict()
        for fila in (
            detalle_rechazados
            .select("campo", "valor", "problema")
            .orderBy("campo", "valor", "problema")
            .limit(20)
            .collect()
        )
    ]

    # Estos indicadores describen conversiones sobre todos los registros.
    # Pueden superponerse; no son el total de rechazos.
    sin_anio_mes = datos.filter(
        F.col("anio").isNull()
        | F.col("mes_numero").isNull()
    ).count()

    sin_dia = datos.filter(
        F.col("dia_semana_numero").isNull()
    ).count()

    sin_kilometro = datos.filter(
        F.col("kilometro_numero").isNull()
    ).count()

    reporte = {
        "perfil": "accidentes_cosevi",
        "registros_leidos": total_leidos,
        "registros_aceptados": total_aceptados,
        "registros_rechazados": total_rechazados,
        "documentos_generados": guardados,
        "puede_importar": total_aceptados > 0,
        "registros_sin_anio_o_mes": sin_anio_mes,
        "registros_sin_dia_semana": sin_dia,
        "registros_sin_kilometro_numerico": sin_kilometro,
        "errores_por_campo": errores_por_campo,
        "muestra_errores": muestra_errores,
        "carpeta_rechazados": (
            "rechazados" if total_rechazados > 0 else None
        ),
        "notas": [
            "Cada registro rechazado informa únicamente su primer fallo.",
            "Solo los registros aceptados generan documentos importables.",
            "El día representa el día de la semana, no el día del mes.",
            "Un kilómetro no numérico no implica descartar el accidente.",
            "Los indicadores de conversión pueden incluir un mismo registro."
        ]
    }

    ruta_reporte = Path(salida) / "reporte.json"
    ruta_reporte.write_text(
        json.dumps(reporte, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

    print("\nRESUMEN FINAL:")
    print("REGISTROS LEÍDOS:", total_leidos)
    print("REGISTROS ACEPTADOS:", total_aceptados)
    print("REGISTROS RECHAZADOS:", total_rechazados)
    print("DOCUMENTOS GUARDADOS:", guardados)
    print("REPORTE GUARDADO EN:", ruta_reporte)

finally:
    try:
        spark.stop()
    finally:
        if temporal is not None:
            temporal.cleanup()
